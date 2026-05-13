from __future__ import annotations

import argparse
import csv
import math
import re
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


BANGKOK_TZ = timezone(timedelta(hours=7))
DEFAULT_SENSOR_CSV = "training-dataset-2026-04-03_to_2026-04-18-cleaned-sensors.csv"
DEFAULT_OUTPUT_PREFIX = "image-features-2026-04-03_to_2026-04-18"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extract OpenCV fixed-mask sky image features and merge them with the cleaned sensor dataset."
    )
    parser.add_argument("--sensor-csv", default=DEFAULT_SENSOR_CSV, help="Cleaned sensor CSV with ImageFileKey column.")
    parser.add_argument("--images-root", default="uploaded-images", help="Root directory that contains saved sky images.")
    parser.add_argument("--output-prefix", default=DEFAULT_OUTPUT_PREFIX, help="Output prefix for generated files.")
    parser.add_argument("--preview-count", type=int, default=12, help="Number of debug preview overlay images to create.")
    parser.add_argument("--analysis-width", type=int, default=360, help="Resize images to this width for faster analysis.")
    return parser.parse_args()


def as_float(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def parse_timestamp_from_file_key(file_key: str) -> tuple[str, str]:
    name = Path(file_key).name
    match = re.match(r"(\d{4}-\d{2}-\d{2})T(\d{2})-(\d{2})-(\d{2})(?:-(\d+))?Z", name)
    if not match:
        return "", ""

    date_part, hh, mm, ss, micros = match.groups()
    microsecond = int((micros or "0")[:6].ljust(6, "0"))
    utc_dt = datetime.strptime(f"{date_part} {hh}:{mm}:{ss}", "%Y-%m-%d %H:%M:%S").replace(
        microsecond=microsecond,
        tzinfo=timezone.utc,
    )
    local_dt = utc_dt.astimezone(BANGKOK_TZ)
    return (
        utc_dt.isoformat().replace("+00:00", "Z"),
        local_dt.strftime("%Y-%m-%d %H:%M:%S"),
    )


def read_sensor_rows(sensor_csv: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    with sensor_csv.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            rows.append(row)
    return rows


def unique_image_keys(sensor_rows: list[dict[str, str]]) -> list[str]:
    keys = {
        row.get("ImageFileKey", "").replace("\\", "/").strip()
        for row in sensor_rows
        if row.get("ImageFileKey")
    }
    return sorted(key for key in keys if key)


def resolve_image_path(images_root: Path, file_key: str) -> Path:
    cleaned = file_key.replace("\\", "/").lstrip("/")
    if cleaned.startswith("uploads/"):
        cleaned = cleaned[len("uploads/") :]
    return images_root / cleaned


def build_fixed_sky_mask(height: int, width: int) -> np.ndarray:
    mask = np.zeros((height, width), dtype=np.uint8)

    # The camera is fixed, so this polygon keeps the sky/lens field and removes
    # most buildings at the bottom plus hard black borders near the frame edges.
    polygon = np.array(
        [
            (int(width * 0.02), int(height * 0.10)),
            (int(width * 0.98), int(height * 0.10)),
            (int(width * 0.99), int(height * 0.78)),
            (int(width * 0.02), int(height * 0.78)),
        ],
        dtype=np.int32,
    )
    cv2.fillPoly(mask, [polygon], 255)

    ellipse = np.zeros_like(mask)
    cv2.ellipse(
        ellipse,
        (int(width * 0.50), int(height * 0.43)),
        (int(width * 0.58), int(height * 0.50)),
        0,
        0,
        360,
        255,
        -1,
    )

    return cv2.bitwise_and(mask, ellipse)


def classify_cloud_status(cloud_coverage: float, quality_score: float) -> str:
    if quality_score < 35:
        return "low_quality"
    if cloud_coverage < 25:
        return "clear_sky"
    if cloud_coverage < 60:
        return "partly_cloudy"
    return "overcast"


def safe_percent(count: int | np.integer, total: int) -> float:
    if total <= 0:
        return 0.0
    return round(float(count) / total * 100, 4)


def analyze_image(
    image_path: Path,
    file_key: str,
    analysis_width: int = 360,
) -> tuple[dict[str, Any], dict[str, np.ndarray] | None]:
    timestamp_utc, timestamp_local = parse_timestamp_from_file_key(file_key)

    base_row: dict[str, Any] = {
        "ImageFileKey": file_key,
        "ImagePath": str(image_path),
        "TimestampUtc": timestamp_utc,
        "TimestampLocal": timestamp_local,
        "ImageReadable": 0,
        "ImageWidth": None,
        "ImageHeight": None,
        "SkyMaskCoveragePercent": None,
        "AnalyzedSkyPixels": 0,
        "CloudCoverageEstimatePercent": None,
        "SkyBrightnessMean": None,
        "SkySaturationMean": None,
        "SkyHueMean": None,
        "BlueSkyPercent": None,
        "WhiteCloudPercent": None,
        "GrayCloudPercent": None,
        "OverexposedPercent": None,
        "DarkPixelPercent": None,
        "SunGlareFlag": 0,
        "BlurScoreLaplacianVar": None,
        "ImageQualityScore": 0,
        "CloudStatusEstimate": "missing_or_unreadable",
    }

    original_image = cv2.imread(str(image_path), cv2.IMREAD_COLOR)
    if original_image is None:
        return base_row, None

    original_height, original_width = original_image.shape[:2]
    image = original_image
    if analysis_width > 0 and original_width > analysis_width:
        scale = analysis_width / original_width
        image = cv2.resize(
            original_image,
            (analysis_width, max(1, int(original_height * scale))),
            interpolation=cv2.INTER_AREA,
        )

    height, width = image.shape[:2]
    hsv = cv2.cvtColor(image, cv2.COLOR_BGR2HSV)
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    blue_channel, green_channel, red_channel = cv2.split(image)

    fixed_mask = build_fixed_sky_mask(height, width)
    h_channel, s_channel, v_channel = cv2.split(hsv)

    # Remove hard lens borders, buildings, and near-black frame regions from the analysis area.
    usable_mask = (
        (fixed_mask > 0)
        & (v_channel > 35)
        & (gray > 30)
    )
    total_fixed_pixels = int(np.count_nonzero(fixed_mask))
    sky_pixels = int(np.count_nonzero(usable_mask))

    if sky_pixels == 0:
        base_row.update(
            {
                "ImageReadable": 1,
                "ImageWidth": original_width,
                "ImageHeight": original_height,
                "SkyMaskCoveragePercent": safe_percent(total_fixed_pixels, width * height),
                "CloudStatusEstimate": "no_sky_pixels",
            }
        )
        return base_row, {"image": image, "fixed_mask": fixed_mask, "usable_mask": usable_mask.astype(np.uint8) * 255}

    sky_h = h_channel[usable_mask]
    sky_s = s_channel[usable_mask]
    sky_v = v_channel[usable_mask]
    sky_gray = gray[usable_mask]

    overexposed_mask = usable_mask & (v_channel >= 245) & (s_channel <= 45)
    dark_mask = (fixed_mask > 0) & (v_channel <= 45)
    blue_dominance = blue_channel.astype(np.int16) - red_channel.astype(np.int16)
    blue_sky_mask = usable_mask & (blue_dominance >= 8) & (v_channel >= 70)

    # Cloud estimate is intentionally conservative: low blue-dominance bright/gray
    # pixels inside the sky mask. This prevents hazy blue sky from becoming 100% cloud.
    white_cloud_mask = usable_mask & (blue_dominance < 8) & (s_channel <= 55) & (v_channel >= 120) & ~overexposed_mask
    gray_cloud_mask = usable_mask & (blue_dominance < 5) & (s_channel <= 80) & (v_channel >= 65) & (v_channel < 120)
    cloud_mask = white_cloud_mask | gray_cloud_mask

    cloud_coverage = safe_percent(np.count_nonzero(cloud_mask), sky_pixels)
    overexposed_percent = safe_percent(np.count_nonzero(overexposed_mask), sky_pixels)
    dark_percent = safe_percent(np.count_nonzero(dark_mask), total_fixed_pixels)
    blue_sky_percent = safe_percent(np.count_nonzero(blue_sky_mask), sky_pixels)
    white_cloud_percent = safe_percent(np.count_nonzero(white_cloud_mask), sky_pixels)
    gray_cloud_percent = safe_percent(np.count_nonzero(gray_cloud_mask), sky_pixels)

    blur_score = float(cv2.Laplacian(sky_gray, cv2.CV_64F).var()) if sky_gray.size else 0.0
    saturation_mean = float(np.mean(sky_s))
    brightness_mean = float(np.mean(sky_v))
    hue_mean = float(np.mean(sky_h))

    overexposed_penalty = min(35.0, overexposed_percent * 1.25)
    dark_penalty = min(60.0, dark_percent * 0.90)
    blur_penalty = 20.0 if blur_score < 12 else 0.0
    low_pixel_penalty = 25.0 if sky_pixels < 10_000 else 0.0
    quality_score = max(0.0, min(100.0, 100.0 - overexposed_penalty - dark_penalty - blur_penalty - low_pixel_penalty))
    sun_glare_flag = 1 if overexposed_percent >= 3.0 else 0

    base_row.update(
        {
            "ImageReadable": 1,
            "ImageWidth": original_width,
            "ImageHeight": original_height,
            "SkyMaskCoveragePercent": safe_percent(total_fixed_pixels, width * height),
            "AnalyzedSkyPixels": sky_pixels,
            "CloudCoverageEstimatePercent": round(cloud_coverage, 4),
            "SkyBrightnessMean": round(brightness_mean, 4),
            "SkySaturationMean": round(saturation_mean, 4),
            "SkyHueMean": round(hue_mean, 4),
            "BlueSkyPercent": blue_sky_percent,
            "WhiteCloudPercent": white_cloud_percent,
            "GrayCloudPercent": gray_cloud_percent,
            "OverexposedPercent": overexposed_percent,
            "DarkPixelPercent": dark_percent,
            "SunGlareFlag": sun_glare_flag,
            "BlurScoreLaplacianVar": round(blur_score, 4),
            "ImageQualityScore": round(quality_score, 4),
            "CloudStatusEstimate": classify_cloud_status(cloud_coverage, quality_score),
        }
    )

    debug_masks = {
        "image": image,
        "fixed_mask": fixed_mask,
        "usable_mask": usable_mask.astype(np.uint8) * 255,
        "cloud_mask": cloud_mask.astype(np.uint8) * 255,
        "overexposed_mask": overexposed_mask.astype(np.uint8) * 255,
    }
    return base_row, debug_masks


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    headers = list(rows[0].keys())
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_excel(path: Path, image_rows: list[dict[str, Any]], summary_rows: list[dict[str, Any]]) -> None:
    workbook = Workbook(write_only=False)
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    write_sheet(summary_sheet, summary_rows)

    image_sheet = workbook.create_sheet("ImageFeatures")
    write_sheet(image_sheet, image_rows)
    workbook.save(path)


def write_sheet(sheet, rows: list[dict[str, Any]]) -> None:
    if not rows:
        sheet["A1"] = "No data"
        return

    headers = list(rows[0].keys())
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    for row in rows:
        sheet.append([row.get(header) for header in headers])

    for column_index, header in enumerate(headers, start=1):
        max_width = min(max(len(str(header)) + 2, 12), 42)
        sheet.column_dimensions[get_column_letter(column_index)].width = max_width


def create_preview(preview_path: Path, masks: dict[str, np.ndarray], row: dict[str, Any]) -> None:
    image = masks["image"].copy()
    cloud_mask = masks.get("cloud_mask")
    overexposed_mask = masks.get("overexposed_mask")
    usable_mask = masks.get("usable_mask")

    overlay = image.copy()
    if usable_mask is not None:
        overlay[usable_mask == 0] = (overlay[usable_mask == 0] * 0.30).astype(np.uint8)
    if cloud_mask is not None:
        overlay[cloud_mask > 0] = (0, 220, 255)
    if overexposed_mask is not None:
        overlay[overexposed_mask > 0] = (0, 0, 255)

    blended = cv2.addWeighted(image, 0.62, overlay, 0.38, 0)
    label = (
        f"cloud={row.get('CloudCoverageEstimatePercent')}% "
        f"quality={row.get('ImageQualityScore')} "
        f"glare={row.get('SunGlareFlag')}"
    )
    cv2.putText(
        blended,
        label,
        (18, 34),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.85,
        (255, 255, 255),
        2,
        cv2.LINE_AA,
    )
    preview_path.parent.mkdir(parents=True, exist_ok=True)
    cv2.imwrite(str(preview_path), blended)


def merge_sensor_image_rows(sensor_rows: list[dict[str, str]], feature_by_key: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    image_feature_columns = [
        "ImageReadable",
        "CloudCoverageEstimatePercent",
        "SkyBrightnessMean",
        "SkySaturationMean",
        "BlueSkyPercent",
        "WhiteCloudPercent",
        "GrayCloudPercent",
        "OverexposedPercent",
        "DarkPixelPercent",
        "SunGlareFlag",
        "BlurScoreLaplacianVar",
        "ImageQualityScore",
        "CloudStatusEstimate",
    ]

    merged: list[dict[str, Any]] = []
    for row in sensor_rows:
        key = row.get("ImageFileKey", "").replace("\\", "/").strip()
        features = feature_by_key.get(key, {})
        merged_row: dict[str, Any] = dict(row)
        for column in image_feature_columns:
            merged_row[column] = features.get(column, "")
        merged.append(merged_row)
    return merged


def write_summary(path: Path, summary: list[dict[str, Any]], status_counts: Counter[str]) -> None:
    lines = [
        "# Image Feature Extraction Summary",
        "",
        "## Scope",
        "- Method: OpenCV fixed sky mask + image feature extraction.",
        "- Cloud coverage is an estimate, not a trained cloud model yet.",
        "- Image processing is separated from the existing sensor-only baseline forecast.",
        "",
        "## Metrics",
        "",
    ]
    for row in summary:
        lines.append(f"- {row['Metric']}: {row['Value']}")

    lines.extend(["", "## Cloud Status Counts", ""])
    for status, count in status_counts.most_common():
        lines.append(f"- {status}: {count}")

    lines.extend(
        [
            "",
            "## Notes For Report",
            "- The fixed mask is suitable because the Raspberry Pi camera angle is fixed.",
            "- Buildings, lens borders, very dark pixels, and overexposed glare are reduced before feature extraction.",
            "- Temperature and humidity remain installation-point sensor values, not true ambient weather.",
            "- Recommended next step: retrain a sensor + image feature model and compare it with the sensor-only baseline.",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    args = parse_args()
    project_root = Path(__file__).resolve().parents[1]
    sensor_csv = Path(args.sensor_csv)
    if not sensor_csv.is_absolute():
        sensor_csv = project_root / sensor_csv

    images_root = Path(args.images_root)
    if not images_root.is_absolute():
        images_root = project_root / images_root

    output_prefix = Path(args.output_prefix)
    if not output_prefix.is_absolute():
        output_prefix = project_root / output_prefix

    image_features_csv = Path(f"{output_prefix}.csv")
    image_features_xlsx = Path(f"{output_prefix}.xlsx")
    merged_sensor_csv = Path(f"{output_prefix}-sensor-image-dataset.csv")
    summary_md = Path(f"{output_prefix}-summary.md")
    preview_dir = Path(f"{output_prefix}-previews")

    sensor_rows = read_sensor_rows(sensor_csv)
    keys = unique_image_keys(sensor_rows)
    preview_key_set: set[str] = set()
    if args.preview_count > 0 and keys:
        if len(keys) <= args.preview_count:
            preview_key_set = set(keys)
        else:
            preview_key_set = {
                keys[round(i * (len(keys) - 1) / (args.preview_count - 1))]
                for i in range(args.preview_count)
            }

    image_rows: list[dict[str, Any]] = []
    feature_by_key: dict[str, dict[str, Any]] = {}
    missing_count = 0
    preview_written = 0

    for index, key in enumerate(keys, start=1):
        image_path = resolve_image_path(images_root, key)
        if not image_path.exists():
            missing_count += 1
        row, masks = analyze_image(image_path, key, analysis_width=args.analysis_width)
        image_rows.append(row)
        feature_by_key[key] = row

        if masks is not None and key in preview_key_set:
            preview_name = f"{preview_written + 1:02d}-{Path(key).stem}.jpg"
            create_preview(preview_dir / preview_name, masks, row)
            preview_written += 1

        if index % 1000 == 0:
            print(f"Processed {index}/{len(keys)} unique images", flush=True)

    merged_rows = merge_sensor_image_rows(sensor_rows, feature_by_key)
    status_counts = Counter(str(row.get("CloudStatusEstimate", "")) for row in image_rows)
    readable_rows = [row for row in image_rows if row.get("ImageReadable") == 1]

    def mean(column: str) -> float:
        values = [as_float(row.get(column), math.nan) for row in readable_rows if row.get(column) not in (None, "")]
        values = [value for value in values if not math.isnan(value)]
        return round(sum(values) / len(values), 4) if values else 0.0

    summary_rows = [
        {"Metric": "Sensor CSV", "Value": str(sensor_csv)},
        {"Metric": "Images root", "Value": str(images_root)},
        {"Metric": "Sensor rows", "Value": len(sensor_rows)},
        {"Metric": "Unique image keys", "Value": len(keys)},
        {"Metric": "Readable images", "Value": len(readable_rows)},
        {"Metric": "Missing/unreadable images", "Value": len(keys) - len(readable_rows)},
        {"Metric": "Missing file paths", "Value": missing_count},
        {"Metric": "Average cloud coverage estimate %", "Value": mean("CloudCoverageEstimatePercent")},
        {"Metric": "Average sky brightness", "Value": mean("SkyBrightnessMean")},
        {"Metric": "Average sky saturation", "Value": mean("SkySaturationMean")},
        {"Metric": "Average overexposed %", "Value": mean("OverexposedPercent")},
        {"Metric": "Average image quality score", "Value": mean("ImageQualityScore")},
        {"Metric": "Preview images created", "Value": preview_written},
    ]

    write_csv(image_features_csv, image_rows)
    write_csv(merged_sensor_csv, merged_rows)
    write_excel(image_features_xlsx, image_rows, summary_rows)
    write_summary(summary_md, summary_rows, status_counts)

    print(f"Image features CSV created: {image_features_csv}")
    print(f"Image features Excel created: {image_features_xlsx}")
    print(f"Sensor + image dataset CSV created: {merged_sensor_csv}")
    print(f"Summary created: {summary_md}")
    print(f"Preview directory: {preview_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
